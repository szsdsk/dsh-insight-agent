"""Independently calculate the pinned UCI Online Retail two-month case."""

from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
EXPECTED = Path(__file__).with_name("expected.json")
DATA = ROOT / "evals" / "runs" / "retail"
ZIP_SHA256 = "f5385cbb54bbebf7196389109c6b0621faab0c304e3702548165e71c84aede8b"
XLSX_SHA256 = "43465a06f2ccf7c8b5bd2892bc7defb52f97487934fe93b16ae4c3936424676d"
MONTHS = {(2011, 9): "baseline", (2011, 10): "current"}


def hash_file(path: Path) -> str:
    """Hash an input file in bounded chunks."""
    digest = sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def change_rows(before: dict[str, Decimal], after: dict[str, Decimal]) -> dict[str, object]:
    """Return the ten largest absolute changes and the complete remainder."""
    ranked = sorted(set(before) | set(after),
                    key=lambda key: (-abs(after.get(key, Decimal(0)) - before.get(key, Decimal(0))), key))
    rows = []
    for key in ranked[:10]:
        old = before.get(key, Decimal(0))
        new = after.get(key, Decimal(0))
        rows.append({"name": key, "baseline": format(old, "f"), "current": format(new, "f"),
                     "delta": format(new - old, "f")})
    remainder = sum((after.get(key, Decimal(0)) - before.get(key, Decimal(0))
                     for key in ranked[10:]), Decimal(0))
    return {"groups": rows, "other_delta": format(remainder, "f"), "group_count": len(ranked)}


def calculate(source: Path, data: Path) -> dict[str, object]:
    """Compute totals directly from the original workbook and pin CSV bytes."""
    if hash_file(source) != ZIP_SHA256:
        raise ValueError("UCI ZIP checksum differs from the pinned download")
    with ZipFile(source) as archive:
        if archive.namelist() != ["Online Retail.xlsx"]:
            raise ValueError("unexpected UCI ZIP contents")
        workbook_bytes = archive.read("Online Retail.xlsx")
    if sha256(workbook_bytes).hexdigest() != XLSX_SHA256:
        raise ValueError("UCI workbook checksum differs from the pinned download")
    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
    if workbook.sheetnames != ["Online Retail"]:
        raise ValueError("unexpected workbook sheets")
    rows = workbook.active.iter_rows(values_only=True)
    if tuple(next(rows)) != ("InvoiceNo", "StockCode", "Description", "Quantity", "InvoiceDate",
                             "UnitPrice", "CustomerID", "Country"):
        raise ValueError("unexpected workbook columns")
    counts = {role: 0 for role in MONTHS.values()}
    totals = {role: Decimal(0) for role in MONTHS.values()}
    cancellations = {role: 0 for role in MONTHS.values()}
    missing_customers = {role: 0 for role in MONTHS.values()}
    countries: dict[str, defaultdict[str, Decimal]] = {role: defaultdict(Decimal) for role in MONTHS.values()}
    products: dict[str, defaultdict[str, Decimal]] = {role: defaultdict(Decimal) for role in MONTHS.values()}
    source_count = 0
    try:
        for invoice, stock, _description, quantity, issued, price, customer, country in rows:
            source_count += 1
            if not isinstance(issued, datetime):
                raise ValueError(f"row {source_count + 1} has no valid InvoiceDate")
            role = MONTHS.get((issued.year, issued.month))
            if role is None:
                continue
            if quantity is None or price is None:
                raise ValueError(f"row {source_count + 1} has no Quantity or UnitPrice")
            amount = Decimal(str(quantity)) * Decimal(str(price))
            counts[role] += 1
            totals[role] += amount
            cancellations[role] += str(invoice).upper().startswith("C")
            missing_customers[role] += customer is None
            countries[role]["" if country is None else str(country)] += amount
            products[role][str(stock)] += amount
    finally:
        workbook.close()
    if source_count != 541_909:
        raise ValueError(f"expected 541909 source rows, found {source_count}")
    periods = {}
    for (year, month), role in MONTHS.items():
        file = data / f"online-retail-{role}.csv"
        periods[role] = {"month": f"{year}-{month:02d}", "rows": counts[role],
                         "signed_revenue": format(totals[role], "f"),
                         "cancellation_rows": cancellations[role],
                         "missing_customer_rows": missing_customers[role],
                         "csv_sha256": hash_file(file)}
    return {"source_zip_sha256": ZIP_SHA256, "source_xlsx_sha256": XLSX_SHA256,
            "source_rows": source_count, "periods": periods,
            "signed_revenue_delta": format(totals["current"] - totals["baseline"], "f"),
            "country_changes": change_rows(countries["baseline"], countries["current"]),
            "product_changes": change_rows(products["baseline"], products["current"])}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True, help="downloaded online+retail.zip")
    parser.add_argument("--data", type=Path, default=DATA)
    parser.add_argument("--print", action="store_true", help="print the independently calculated values")
    arguments = parser.parse_args()
    actual = calculate(arguments.source, arguments.data)
    if arguments.print:
        print(json.dumps(actual, indent=2, ensure_ascii=False))
    elif actual != json.loads(EXPECTED.read_text(encoding="utf-8")):
        raise SystemExit("UCI Online Retail results differ from evals/retail/expected.json")
    else:
        print("UCI Online Retail source, extract, and independent oracle agree")

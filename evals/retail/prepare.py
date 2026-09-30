"""Extract two fixed months from the pinned UCI Online Retail workbook."""

from __future__ import annotations

import argparse
import csv
from datetime import datetime
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "evals" / "runs" / "retail"
ZIP_SHA256 = "f5385cbb54bbebf7196389109c6b0621faab0c304e3702548165e71c84aede8b"
XLSX_SHA256 = "43465a06f2ccf7c8b5bd2892bc7defb52f97487934fe93b16ae4c3936424676d"
SOURCE_ROWS = 541_909
SOURCE_COLUMNS = ("InvoiceNo", "StockCode", "Description", "Quantity", "InvoiceDate",
                  "UnitPrice", "CustomerID", "Country")
OUTPUT_COLUMNS = ("InvoiceNo", "StockCode", "InvoiceDate", "Quantity", "UnitPrice",
                  "CustomerID", "Country", "LineRevenue")
PERIODS = {(2011, 9): "baseline", (2011, 10): "current"}


def digest(path: Path) -> str:
    """Return a file's SHA-256 without loading it into memory."""
    value = sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            value.update(block)
    return value.hexdigest()


def extract(source: Path, output: Path) -> dict[str, object]:
    """Select September and October 2011 lines with signed line revenue."""
    if digest(source) != ZIP_SHA256:
        raise ValueError("UCI ZIP checksum differs from the pinned download")
    with ZipFile(source) as archive:
        if archive.namelist() != ["Online Retail.xlsx"]:
            raise ValueError("unexpected UCI ZIP contents")
        workbook_bytes = archive.read("Online Retail.xlsx")
    if sha256(workbook_bytes).hexdigest() != XLSX_SHA256:
        raise ValueError("UCI workbook checksum differs from the pinned download")
    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
    if workbook.sheetnames != ["Online Retail"]:
        raise ValueError("unexpected workbook sheets")
    sheet = workbook["Online Retail"]
    rows = sheet.iter_rows(values_only=True)
    if tuple(next(rows)) != SOURCE_COLUMNS:
        raise ValueError("unexpected workbook columns")
    output.mkdir(parents=True, exist_ok=True)
    paths = {role: output / f"online-retail-{role}.csv" for role in PERIODS.values()}
    handles = {role: path.with_suffix(".partial").open("w", encoding="utf-8", newline="")
               for role, path in paths.items()}
    counts = {role: 0 for role in paths}
    try:
        writers = {role: csv.writer(handle, lineterminator="\n") for role, handle in handles.items()}
        for writer in writers.values():
            writer.writerow(OUTPUT_COLUMNS)
        source_count = 0
        for row in rows:
            source_count += 1
            invoice, stock, _description, quantity, issued, price, customer, country = row
            if not isinstance(issued, datetime):
                raise ValueError(f"row {source_count + 1} has no valid InvoiceDate")
            role = PERIODS.get((issued.year, issued.month))
            if role is None:
                continue
            if quantity is None or price is None:
                raise ValueError(f"row {source_count + 1} has no Quantity or UnitPrice")
            revenue = Decimal(str(quantity)) * Decimal(str(price))
            writers[role].writerow((str(invoice), str(stock), issued.isoformat(sep=" "),
                                    str(quantity), str(price), "" if customer is None else str(int(customer)),
                                    "" if country is None else str(country), format(revenue, "f")))
            counts[role] += 1
    finally:
        workbook.close()
        for handle in handles.values():
            handle.close()
    if source_count != SOURCE_ROWS:
        raise ValueError(f"expected {SOURCE_ROWS} source rows, found {source_count}")
    for role, path in paths.items():
        path.with_suffix(".partial").replace(path)
    return {"source_zip_sha256": ZIP_SHA256, "source_xlsx_sha256": XLSX_SHA256,
            "source_rows": source_count,
            "periods": {role: {"month": f"2011-{month:02d}", "rows": counts[role],
                               "file": path.name, "sha256": digest(path)}
                        for (year, month), role in PERIODS.items() for path in [paths[role]]}}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True, help="downloaded online+retail.zip")
    parser.add_argument("--output", type=Path, default=OUTPUT)
    arguments = parser.parse_args()
    print(json.dumps(extract(arguments.source, arguments.output), indent=2, ensure_ascii=False))

from __future__ import annotations

from contextlib import closing
from decimal import Decimal
import json
import re
import threading
import time
from typing import Any, TYPE_CHECKING

import duckdb
from openpyxl.utils import get_column_letter

from .models import ComparisonSpec, QueryRecord, QueryResult, TableRef
from .serialization import to_json_value
from .sql_policy import validate_read_only_sql

if TYPE_CHECKING:
    from .engine import DataEngine


def diagnose_table(
    engine: DataEngine,
    source_id: str,
    relation: str,
    key_columns: list[str] | None = None,
    metric_columns: list[str] | None = None,
) -> dict[str, Any]:
    from .engine import SourceError, quote_identifier, quote_literal

    source = engine._source(source_id)
    engine._require_relation(source, relation)
    schema = engine.describe_relation(source_id, relation)["columns"]
    names = [column["name"] for column in schema]
    keys = key_columns or []
    metrics = metric_columns or []
    if not names:
        raise SourceError("selected table has no columns")
    if any(name not in names for name in keys + metrics):
        raise SourceError("diagnostic field is absent from the selected table")
    if len(set(keys)) != len(keys) or len(set(metrics)) != len(metrics):
        raise SourceError("diagnostic fields must be unique")
    if len(keys) > 4 or len(metrics) > 3:
        raise SourceError("diagnostic key or metric limit exceeded")
    table = quote_identifier(relation)
    ordinal = "__insight_position"
    while ordinal in names:
        ordinal += "_"
    position = quote_identifier(ordinal)
    sampled: list[tuple[str, list[str], str]] = []
    def sample_projection(key: str, fields: list[str], predicate: str) -> str:
        sampled.append((key, fields, predicate))
        if source.kind == "sqlite":
            return f"'[]' AS {key}_samples"
        displayed = fields[:8]
        pairs = ", ".join(
            f"{quote_literal(name)}, LEFT(CAST({quote_identifier(name)} AS VARCHAR), 200)"
            for name in displayed
        )
        return (f"(SELECT COALESCE(to_json(list(sample ORDER BY sample_position)), '[]') FROM "
                f"(SELECT {position} AS sample_position, json_object('table_row', {position}, "
                f"'values', json_object({pairs}), 'values_truncated', "
                f"({'TRUE' if len(fields) > 8 else 'FALSE'} OR "
                f"COALESCE(({' OR '.join(f'LENGTH(CAST({quote_identifier(name)} AS VARCHAR)) > 200' for name in displayed)}), FALSE))) AS sample "
                f"FROM numbered WHERE {predicate} ORDER BY {position} "
                f"LIMIT {engine.diagnostic_sample_limit})) AS {key}_samples")
    projections = ["COUNT(*) AS row_count"]
    for index, name in enumerate(names):
        field = quote_identifier(name)
        projections.append(
            f"COUNT(*) - COUNT({field}) + "
            f"COUNT(*) FILTER (WHERE TRIM(CAST({field} AS VARCHAR)) = '') AS missing_{index}"
        )
        projections.append(sample_projection(f"missing_{index}", [name],
            f"{field} IS NULL OR TRIM(CAST({field} AS VARCHAR)) = ''"))
    projections.append(
        f"COUNT(*) - (SELECT COUNT(*) FROM (SELECT DISTINCT * FROM {table})) AS duplicate_rows"
    )
    def duplicate_condition(fields: list[str]) -> str:
        return " AND ".join(f"peer.{quote_identifier(name)} IS NOT DISTINCT FROM numbered.{quote_identifier(name)}" for name in fields)
    projections.append(sample_projection("duplicate_rows", names,
        f"EXISTS (SELECT 1 FROM numbered peer WHERE peer.{position} <> numbered.{position} AND {duplicate_condition(names)})"))
    if keys:
        quoted = ", ".join(quote_identifier(name) for name in keys)
        projections.append(
            f"COUNT(*) - (SELECT COUNT(*) FROM (SELECT DISTINCT {quoted} FROM {table})) AS duplicate_keys"
        )
        projections.append(sample_projection("duplicate_keys", keys,
            f"EXISTS (SELECT 1 FROM numbered peer WHERE peer.{position} <> numbered.{position} AND {duplicate_condition(keys)})"))
    for index, name in enumerate(metrics):
        field = quote_identifier(name)
        projections.append(
            f"COUNT(*) FILTER (WHERE {field} IS NOT NULL AND "
            f"TRIM(CAST({field} AS VARCHAR)) <> '' AND "
            f"TRY_CAST({field} AS DOUBLE) IS NULL) AS invalid_numeric_{index}"
        )
        projections.append(sample_projection(f"invalid_numeric_{index}", [name],
            f"{field} IS NOT NULL AND TRIM(CAST({field} AS VARCHAR)) <> '' AND TRY_CAST({field} AS DOUBLE) IS NULL"))
    sql = f"WITH numbered AS (SELECT *, ROW_NUMBER() OVER () AS {position} FROM {table}) SELECT {', '.join(projections)} FROM {table}"
    result = engine.execute_sql(source_id, sql, 1)
    values = dict(zip(result["columns"], result["rows"][0], strict=True))
    total = values["row_count"]
    locations: dict[str, Any] | None = None
    path = engine._source_path(source)
    if source.kind == "xlsx":
        sheets, _ = engine._read_xlsx(path, source.selection)
        locations = sheets[relation]
    elif source.kind == "csv" and source.selection is not None:
        locations = engine._read_csv(path, source.selection)
    row_numbers = locations["row_numbers"] if locations is not None and len(locations["row_numbers"]) == total else []
    def locate(sample: dict[str, Any], fields: list[str]) -> dict[str, Any]:
        if row_numbers:
            sample["source_row"] = row_numbers[sample["table_row"] - 1]
        if source.kind == "xlsx":
            sample["sheet"] = relation
            if "source_row" in sample:
                sample["cells"] = [f"{get_column_letter(names.index(name) + 1)}{sample['source_row']}" for name in fields[:8]]
        return sample
    samples = {key: [locate(sample, fields) for sample in json.loads(values[f"{key}_samples"])]
               for key, fields, _ in sampled}
    findings = []
    for index, name in enumerate(names):
        count = values[f"missing_{index}"]
        if count:
            findings.append({"kind": "missing", "field": name, "count": count, "rate": count / total,
                             "classification": "review" if name in metrics else "observation",
                             "query_id": result["query_id"], "samples": samples[f"missing_{index}"]})
    for kind, key in (("duplicate_rows", "duplicate_rows"), ("duplicate_keys", "duplicate_keys")):
        if key in values and values[key]:
            findings.append({"kind": kind, "fields": keys if kind == "duplicate_keys" else names, "count": values[key], "rate": values[key] / total,
                             "classification": "review" if kind == "duplicate_keys" else "observation",
                             "query_id": result["query_id"], "samples": samples[key]})
    for index, name in enumerate(metrics):
        count = values[f"invalid_numeric_{index}"]
        if count:
            findings.append({"kind": "invalid_numeric", "field": name, "count": count, "rate": count / total,
                             "classification": "blocking", "query_id": result["query_id"],
                             "samples": samples[f"invalid_numeric_{index}"]})
    warnings = engine._source_warnings.get(source_id, [])
    for warning in warnings:
        match = re.search(r"(\d+) formula cells have no cached value", warning)
        if match and warning.startswith(f"{relation}:"):
            count = int(match.group(1))
            formula_samples = []
            for entry in (locations or {}).get("formula_locations", []):
                column = entry["column_index"]
                formula_samples.append({"source_row": entry["source_row"], "sheet": relation,
                                        "cells": [f"{get_column_letter(column + 1)}{entry['source_row']}"],
                                        "values": entry["values"]})
            findings.append({"kind": "formula_cache_missing", "count": count,
                             "rate": count / max(1, total * len(names)), "classification": "review",
                             "samples": formula_samples})
    return {**result, "row_total": total, "findings": findings, "source_warnings": warnings}


def compare_tables(engine: DataEngine, spec: ComparisonSpec) -> dict[str, Any]:
    from .engine import SourceError

    if spec.baseline.source_id == spec.current.source_id and spec.baseline.relation == spec.current.relation:
        raise SourceError("comparison requires distinct selected tables")
    names = [column.name for column in spec.columns]
    if len(set(names)) != len(names) or any(not name.strip() for name in names):
        raise SourceError("comparison mapping names must be non-empty and unique")
    mapped = set(names)
    if any(name not in mapped for name in spec.dimensions):
        raise SourceError("comparison dimension is not mapped")
    if len(set(spec.dimensions)) != len(spec.dimensions):
        raise SourceError("comparison dimensions must be unique")
    if len(set(metric.name for metric in spec.metrics)) != len(spec.metrics):
        raise SourceError("comparison metric names must be unique")
    for metric in spec.metrics:
        if metric.field is None and metric.aggregation != "count":
            raise SourceError(f"metric {metric.name} requires a field")
        if metric.field is not None and metric.field not in mapped:
            raise SourceError(f"metric field is not mapped: {metric.field}")
    if any(item.field not in mapped for item in spec.filters):
        raise SourceError("comparison filter field is not mapped")

    sources = [engine._source(spec.baseline.source_id), engine._source(spec.current.source_id)]
    schemas = []
    for ref in (spec.baseline, spec.current):
        engine._require_relation(engine._source(ref.source_id), ref.relation)
        schemas.append({item["name"]: item["type"] for item in engine.describe_relation(ref.source_id, ref.relation)["columns"]})
    for column in spec.columns:
        if column.baseline not in schemas[0] or column.current not in schemas[1]:
            raise SourceError(f"mapped column is absent: {column.name}")
        for metric in spec.metrics:
            if metric.field == column.name and metric.aggregation in ("sum", "avg"):
                types = (schemas[0][column.baseline], schemas[1][column.current])
                if not all(_numeric_type(kind) for kind in types):
                    raise SourceError(f"metric {metric.name} has incompatible numeric columns")

    connection = duckdb.connect(":memory:")
    started = time.perf_counter()
    try:
        for role, source, ref, schema in zip(("baseline", "current"), sources, (spec.baseline, spec.current), schemas, strict=True):
            _copy_selection(engine, connection, role, source, ref, schema, spec)
        connection.execute("SET enable_external_access = false")
        result = _comparison_query(engine, connection, spec, sources, started)
        if result["truncated"]:
            raise SourceError("comparison has too many groups; choose a filter or another dimension")
        rows = result["rows"]
        dimensions = len(spec.dimensions)
        if dimensions:
            total_rows = [row for row in rows if row[0] == (1 << dimensions) - 1]
            if len(total_rows) != 1:
                raise SourceError("comparison total is absent")
            total_row = total_rows[0]
            group_rows = [row for row in rows if row[0] == 0]
        else:
            total_row = rows[0]
            group_rows = []
        offset = 1 + dimensions
        totals = [
            {"name": metric.name, **_metric_change(total_row[offset + index * 2], total_row[offset + index * 2 + 1])}
            for index, metric in enumerate(spec.metrics)
        ]
        metric_summaries = []
        for index, metric in enumerate(spec.metrics):
            values = [_metric_change(row[offset + index * 2], row[offset + index * 2 + 1]) for row in group_rows]
            metric_summaries.append({"name": metric.name, "groups": values})
        groups = [{"dimensions": row[1:offset], "metrics": [summary["groups"][index] for summary in metric_summaries]}
                  for index, row in enumerate(group_rows)]
        group_count = int(total_row[-1]) if dimensions else 0
        if group_count > len(group_rows):
            groups.append({"dimensions": ["Other"] * dimensions, "is_other": True,
                           "combined_groups": group_count - len(group_rows), "metrics": [
                _other_metric(totals[metric_index], summary["groups"], spec.metrics[metric_index].aggregation)
                for metric_index, summary in enumerate(metric_summaries)
            ]})
        for group in groups:
            for index, item in enumerate(group["metrics"]):
                delta = item["delta"]
                total_delta = totals[index]["delta"]
                item["contribution_rate"] = (
                    delta / total_delta
                    if spec.metrics[index].aggregation in ("sum", "count")
                    and delta is not None and total_delta not in (None, 0)
                    else None
                )
        return {**result, "metrics": [metric.model_dump() for metric in spec.metrics], "totals": totals,
                "groups": groups, "group_count": group_count,
                "mapping": [column.model_dump() for column in spec.columns]}
    finally:
        connection.close()


def _copy_selection(
    engine: DataEngine,
    target: duckdb.DuckDBPyConnection,
    role: str,
    source: Any,
    ref: TableRef,
    schema: dict[str, str],
    spec: ComparisonSpec,
) -> None:
    from .engine import QueryTimeoutError, SourceError, quote_identifier, quote_literal

    mappings = [(column.name, column.baseline if role == "baseline" else column.current) for column in spec.columns]
    if source.kind == "csv" and source.selection is None:
        path = engine._source_path(source)
        projections = ", ".join(
            f"CAST({quote_identifier(field)} AS {_comparison_type(schema[field])}) AS {quote_identifier(name)}"
            for name, field in mappings
        )
        timer = threading.Timer(engine.timeout_seconds, target.interrupt)
        timer.daemon = True
        timer.start()
        try:
            target.execute(
                f"CREATE TABLE {role}_input AS SELECT {projections} "
                f"FROM read_csv_auto({quote_literal(str(path))}, sample_size = -1)"
            )
            copied = target.execute(f"SELECT COUNT(*) FROM {role}_input").fetchone()[0] * len(mappings)
        except duckdb.Error as error:
            if "interrupt" in str(error).lower():
                raise QueryTimeoutError(f"comparison input exceeded {engine.timeout_seconds:g} seconds") from error
            raise
        finally:
            timer.cancel()
        if copied > engine.max_xlsx_cells:
            raise SourceError("comparison input exceeds the configured cell limit")
        engine._source_path(source)
        return
    definitions = ", ".join(
        f"{quote_identifier(name)} {_comparison_type(schema[field])}" for name, field in mappings
    )
    target.execute(f"CREATE TABLE {role}_input ({definitions})")
    fields = ", ".join(quote_identifier(field) for _, field in mappings)
    sql = f"SELECT {fields} FROM {quote_identifier(ref.relation)}"
    with closing(engine._connect_sqlite(source) if source.kind == "sqlite" else engine._connect_duckdb(source)) as origin:
        cursor = origin.execute(sql)
        copied = 0
        placeholders = ", ".join("?" for _ in mappings)
        while batch := cursor.fetchmany(1_000):
            copied += len(batch) * len(mappings)
            if copied > engine.max_xlsx_cells:
                raise SourceError("comparison input exceeds the configured cell limit")
            target.executemany(f"INSERT INTO {role}_input VALUES ({placeholders})", batch)
    engine._source_path(source)


def _comparison_query(
    engine: DataEngine,
    connection: duckdb.DuckDBPyConnection,
    spec: ComparisonSpec,
    sources: list[Any],
    started: float,
) -> dict[str, Any]:
    from .engine import QueryTimeoutError, SourceError, bind_sql_parameters, quote_identifier

    dimensions = [quote_identifier(name) for name in spec.dimensions]
    predicates = []
    parameters = []
    operators = {"eq": "=", "ne": "<>", "gt": ">", "gte": ">=", "lt": "<", "lte": "<="}
    for item in spec.filters:
        field = quote_identifier(item.field)
        if item.operator in operators:
            predicates.append(f"{field} {operators[item.operator]} ?")
            parameters.append(item.value)
        elif item.operator == "is_null":
            predicates.append(f"{field} IS NULL")
        else:
            predicates.append(f"{field} IS NOT NULL")
    where = " WHERE " + " AND ".join(predicates) if predicates else ""
    aggregations = []
    for index, metric in enumerate(spec.metrics):
        field = quote_identifier(metric.field) if metric.field is not None else "*"
        func = {"sum": "SUM", "count": "COUNT", "count_distinct": "COUNT", "avg": "AVG"}[metric.aggregation]
        arg = f"DISTINCT {field}" if metric.aggregation == "count_distinct" else field
        calculation = f"{func}({arg})"
        if metric.aggregation == "sum":
            calculation = f"COALESCE({calculation}, 0)"
        aggregations.append(f"{calculation} AS m{index}")
    group_by = " GROUP BY GROUPING SETS ((" + ", ".join(dimensions) + "), ())" if dimensions else ""
    projections = ", ".join(dimensions + ["COUNT(*) AS group_rows"] + aggregations)
    scope = f"GROUPING_ID({', '.join(dimensions)})" if dimensions else "0"
    ctes = [f"{role} AS (SELECT {scope} AS rollup_level, {projections} FROM {role}_input{where}{group_by})" for role in ("baseline", "current")]
    select = ["COALESCE(b.rollup_level, c.rollup_level) AS rollup_level"] + [f"COALESCE(b.{name}, c.{name}) AS {name}" for name in dimensions]
    for index, metric in enumerate(spec.metrics):
        base = f"b.m{index}"
        current = f"c.m{index}"
        if metric.aggregation in ("sum", "count", "count_distinct"):
            base = f"CASE WHEN b.group_rows IS NULL THEN 0 ELSE {base} END"
            current = f"CASE WHEN c.group_rows IS NULL THEN 0 ELSE {current} END"
        select.extend((f"{base} AS baseline_{index}", f"{current} AS current_{index}"))
    join = (
        "FULL OUTER JOIN current c ON b.rollup_level = c.rollup_level AND "
        + " AND ".join(f"b.{name} IS NOT DISTINCT FROM c.{name}" for name in dimensions)
        if dimensions else "CROSS JOIN current c"
    )
    joined = f"SELECT {', '.join(select)} FROM baseline b {join}"
    if dimensions:
        total_level = (1 << len(dimensions)) - 1
        output_columns = ["rollup_level", *dimensions]
        for index in range(len(spec.metrics)):
            output_columns.extend((f"baseline_{index}", f"current_{index}"))
        output_columns.append("comparison_group_count")
        ranking = f"ABS(COALESCE(current_0 - baseline_0, 0)) DESC, {', '.join(dimensions)}"
        sql = (
            f"WITH {', '.join(ctes)}, joined AS ({joined}), ranked AS ("
            f"SELECT *, ROW_NUMBER() OVER (PARTITION BY rollup_level ORDER BY {ranking}) AS group_rank, "
            "COUNT(*) FILTER (WHERE rollup_level = 0) OVER () AS comparison_group_count FROM joined) "
            f"SELECT {', '.join(output_columns)} FROM ranked "
            f"WHERE rollup_level = {total_level} OR (rollup_level = 0 AND group_rank <= {spec.top_n}) "
            "ORDER BY rollup_level, group_rank"
        )
    else:
        sql = f"WITH {', '.join(ctes)} {joined}"
    executable_sql = sql
    display_sql = bind_sql_parameters(executable_sql, parameters * 2)
    validated = validate_read_only_sql(display_sql, dialect="duckdb")
    timer = threading.Timer(engine.timeout_seconds, connection.interrupt)
    timer.daemon = True
    timer.start()
    try:
        cursor = connection.execute(executable_sql, parameters * 2)
        columns = [item[0] for item in cursor.description or []]
        fetched = cursor.fetchmany(engine.default_max_rows + 1)
    except duckdb.Error as error:
        if "interrupt" in str(error).lower():
            raise QueryTimeoutError(f"comparison exceeded {engine.timeout_seconds:g} seconds") from error
        raise
    finally:
        timer.cancel()
    if len(fetched) > engine.default_max_rows:
        raise SourceError("comparison has too many groups; choose a filter or another dimension")
    for source in sources:
        engine._source_path(source)
    rows = [to_json_value(row) for row in fetched]
    digest = "sha256:" + engine._digest({"columns": columns, "rows": rows})
    source_id = "pair_" + engine._digest([source.source_id for source in sources])[:16]
    fingerprint = engine._digest([source.fingerprint for source in sources])
    query_id = "qry_" + engine._digest({"source_id": source_id, "sql": validated.normalized, "result_digest": digest})[:20]
    result = QueryResult(
        query_id=query_id, source_id=source_id, source_fingerprint=fingerprint,
        sql=validated.normalized, columns=columns, rows=rows, row_count=len(rows),
        truncated=False, elapsed_ms=max(0, round((time.perf_counter() - started) * 1_000)),
        result_digest=digest, sources=sources,
    )
    with engine._lock:
        engine._queries[query_id] = QueryRecord(result=result)
        engine._verified_queries.discard(query_id)
    return result.model_dump(mode="json")


def _numeric_type(value: str) -> bool:
    upper = value.upper()
    return any(token in upper for token in ("INT", "DOUBLE", "FLOAT", "DECIMAL", "NUMERIC", "REAL"))


def _comparison_type(value: str) -> str:
    return "DOUBLE" if _numeric_type(value) else "VARCHAR"


def _metric_change(baseline: Any, current: Any) -> dict[str, Any]:
    baseline = to_json_value(baseline)
    current = to_json_value(current)
    if baseline is None or current is None:
        return {"baseline": baseline, "current": current, "delta": None, "change_rate": None}
    delta = Decimal(str(current)) - Decimal(str(baseline))
    rate = delta / Decimal(str(baseline)) if Decimal(str(baseline)) > 0 else None
    return {"baseline": baseline, "current": current, "delta": float(delta), "change_rate": float(rate) if rate is not None else None}


def _other_metric(total: dict[str, Any], kept: list[dict[str, Any]], kind: str) -> dict[str, Any]:
    if kind not in ("sum", "count"):
        return {"baseline": None, "current": None, "delta": None, "change_rate": None}
    baseline = Decimal(str(total["baseline"] or 0)) - sum(
        (Decimal(str(value["baseline"] or 0)) for value in kept), Decimal(0))
    current = Decimal(str(total["current"] or 0)) - sum(
        (Decimal(str(value["current"] or 0)) for value in kept), Decimal(0))
    return _metric_change(baseline, current)

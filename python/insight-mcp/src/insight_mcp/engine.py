from __future__ import annotations

from contextlib import closing
import csv
import hashlib
import os
import re
import sqlite3
import threading
import time
from datetime import date, datetime
from decimal import Decimal
from itertools import islice, zip_longest
from pathlib import Path
from typing import Any, Sequence
from urllib.parse import quote

import duckdb
from openpyxl import load_workbook

from .models import AnalysisSpec, ComparisonSpec, FieldRef, QueryRecord, QueryResult, SourceInfo, SourceKind, TableSelection
from .security import resolve_workspace_path, workspace_relative_path
from .discovery import list_source_files as discover_source_files
from .serialization import canonical_json, to_json_value
from .sql_policy import SqlPolicyError, validate_read_only_sql


class SourceError(ValueError):
    pass


class QueryTimeoutError(TimeoutError):
    pass


class DataEngine:
    HARD_MAX_ROWS = 5_000
    HARD_MAX_TIMEOUT_SECONDS = 30.0
    CSV_RELATION = "data"
    DEFAULT_MAX_XLSX_BYTES = 50 * 1024 * 1024
    DEFAULT_MAX_XLSX_CELLS = 1_000_000

    def __init__(
        self,
        workspace_root: Path | str,
        *,
        default_max_rows: int = 200,
        timeout_seconds: float = 10.0,
        max_xlsx_bytes: int = DEFAULT_MAX_XLSX_BYTES,
        max_xlsx_cells: int = DEFAULT_MAX_XLSX_CELLS,
    ) -> None:
        self.workspace_root = Path(workspace_root).resolve(strict=True)
        self.default_max_rows = self._bounded_rows(default_max_rows)
        if not 0 < timeout_seconds <= self.HARD_MAX_TIMEOUT_SECONDS:
            raise ValueError("timeout_seconds must be within (0, 30]")
        self.timeout_seconds = timeout_seconds
        self.max_xlsx_bytes = int(max_xlsx_bytes)
        self.max_xlsx_cells = int(max_xlsx_cells)
        if self.max_xlsx_bytes <= 0 or self.max_xlsx_cells <= 0:
            raise ValueError("Excel limits must be positive")
        self._sources: dict[str, SourceInfo] = {}
        self._queries: dict[str, QueryRecord] = {}
        self._verified_queries: set[str] = set()
        self._source_warnings: dict[str, list[str]] = {}
        self._lock = threading.RLock()

    @classmethod
    def from_environment(cls) -> "DataEngine":
        root = Path(os.environ.get("INSIGHT_WORKSPACE", os.getcwd()))
        max_rows = int(os.environ.get("INSIGHT_MAX_ROWS", "200"))
        timeout = float(os.environ.get("INSIGHT_QUERY_TIMEOUT_SECONDS", "10"))
        max_xlsx_bytes = int(os.environ.get("INSIGHT_MAX_XLSX_BYTES", str(cls.DEFAULT_MAX_XLSX_BYTES)))
        max_xlsx_cells = int(os.environ.get("INSIGHT_MAX_XLSX_CELLS", str(cls.DEFAULT_MAX_XLSX_CELLS)))
        return cls(
            root,
            default_max_rows=max_rows,
            timeout_seconds=timeout,
            max_xlsx_bytes=max_xlsx_bytes,
            max_xlsx_cells=max_xlsx_cells,
        )

    def list_source_files(self, directory: str = ".") -> dict[str, Any]:
        """Discover workspace data files within bounded traversal limits."""
        return discover_source_files(self.workspace_root, directory)

    def register_source(
        self, path: str, kind: SourceKind, selection: TableSelection | None = None
    ) -> dict[str, Any]:
        resolved = resolve_workspace_path(self.workspace_root, path)
        detected = self._validate_kind(resolved, kind)
        if selection is not None:
            if kind not in ("csv", "xlsx"):
                raise SourceError("table selection requires CSV or XLSX")
            if kind == "csv" and selection.sheet is not None:
                raise SourceError("CSV table selection cannot name a sheet")
            if selection.data_start_row is not None and selection.data_start_row <= selection.header_row:
                raise SourceError("data_start_row must follow header_row")
            if selection.data_end_row is not None and selection.data_end_row < (selection.data_start_row or selection.header_row + 1):
                raise SourceError("data_end_row must include the first data row")
        if detected == "xlsx" and resolved.stat().st_size > self.max_xlsx_bytes:
            raise SourceError(
                f"Excel file exceeds {self.max_xlsx_bytes} byte limit: {resolved.name}"
            )
        relative = workspace_relative_path(self.workspace_root, resolved)
        fingerprint = self._file_fingerprint(resolved)
        source_id = "src_" + self._digest(
            {"kind": detected, "path": relative, "fingerprint": fingerprint,
             "selection": selection.model_dump() if selection is not None else None}
        )[:16]
        info = SourceInfo(
            source_id=source_id,
            kind=detected,
            path=relative,
            fingerprint=fingerprint,
            selection=selection,
        )
        warnings: list[str] = []
        if detected == "xlsx":
            _, warnings = self._read_xlsx(resolved, selection)
        elif detected == "csv" and selection is not None:
            self._read_csv(resolved, selection)
        with self._lock:
            self._sources[source_id] = info
            self._source_warnings[source_id] = warnings
        return {**info.model_dump(mode="json"), "warnings": warnings}

    def preview_table(self, path: str, kind: SourceKind, sheet: str | None = None) -> dict[str, Any]:
        resolved = resolve_workspace_path(self.workspace_root, path)
        self._validate_kind(resolved, kind)
        if kind not in ("csv", "xlsx"):
            raise SourceError("table preview requires CSV or XLSX")
        if resolved.stat().st_size > self.max_xlsx_bytes:
            raise SourceError("preview file exceeds the configured byte limit")
        if kind == "csv":
            if sheet is not None:
                raise SourceError("CSV preview cannot name a sheet")
            with resolved.open(encoding="utf-8-sig", newline="") as stream:
                rows = list(islice(csv.reader(stream), 12))
            return {"sheets": [], "sheet": None, "rows": [row[:32] for row in rows]}
        with closing(load_workbook(resolved, read_only=True, data_only=True, keep_links=False)) as workbook:
            selected = sheet or workbook.sheetnames[0]
            if selected not in workbook.sheetnames:
                raise SourceError(f"unknown sheet: {selected}")
            rows = [
                [to_json_value(value) for value in row[:32]]
                for row in islice(workbook[selected].iter_rows(values_only=True), 12)
            ]
            return {"sheets": workbook.sheetnames, "sheet": selected, "rows": rows}

    def list_relations(self, source_id: str) -> dict[str, Any]:
        source = self._source(source_id)
        # Even metadata-only operations must detect a file that was replaced
        # by an escaping symlink after registration.
        self._source_path(source)
        if source.kind == "csv":
            relations = [self.CSV_RELATION]
        elif source.kind == "xlsx":
            if source.selection is not None and source.selection.sheet is not None:
                relations = [source.selection.sheet]
            else:
                workbook = load_workbook(
                    self._source_path(source), read_only=True, data_only=True, keep_links=False
                )
                try:
                    relations = list(workbook.sheetnames)
                finally:
                    workbook.close()
        elif source.kind == "sqlite":
            with closing(self._connect_sqlite(source)) as connection:
                rows = connection.execute(
                    "SELECT name FROM sqlite_master "
                    "WHERE type IN ('table', 'view') AND name NOT LIKE 'sqlite_%' ORDER BY name"
                ).fetchall()
            relations = [str(row[0]) for row in rows]
        else:
            with closing(self._connect_duckdb(source)) as connection:
                rows = connection.execute(
                    "SELECT table_name FROM information_schema.tables "
                    "WHERE table_schema NOT IN ('information_schema', 'pg_catalog') "
                    "ORDER BY table_name"
                ).fetchall()
            relations = [str(row[0]) for row in rows]
        return {
            "source_id": source_id,
            "relations": relations,
            "warnings": list(self._source_warnings.get(source_id, [])),
        }

    def describe_relation(self, source_id: str, relation: str) -> dict[str, Any]:
        source = self._source(source_id)
        self._require_relation(source, relation)
        if source.kind == "sqlite":
            with closing(self._connect_sqlite(source)) as connection:
                rows = connection.execute(
                    f"PRAGMA table_info({quote_identifier(relation)})"
                ).fetchall()
            columns = [
                {
                    "name": str(row[1]),
                    "type": str(row[2] or "UNKNOWN"),
                    "nullable": not bool(row[3]),
                    "primary_key": bool(row[5]),
                }
                for row in rows
            ]
        else:
            with closing(self._connect_duckdb(source)) as connection:
                cursor = connection.execute(
                    f"DESCRIBE SELECT * FROM {quote_identifier(relation)}"
                )
                rows = cursor.fetchall()
            columns = [
                {
                    "name": str(row[0]),
                    "type": str(row[1]),
                    "nullable": str(row[2]).upper() != "NO",
                    "primary_key": False,
                }
                for row in rows
            ]
        return {
            "source_id": source_id,
            "relation": relation,
            "columns": columns,
            "warnings": list(self._source_warnings.get(source_id, [])),
        }

    def profile_relation(
        self,
        source_id: str,
        relation: str,
        columns: Sequence[str] | None = None,
    ) -> dict[str, Any]:
        source = self._source(source_id)
        schema = self.describe_relation(source_id, relation)["columns"]
        available = {item["name"]: item["type"] for item in schema}
        selected = list(columns) if columns else list(available)
        if len(selected) > 32:
            raise SourceError("profile_relation accepts at most 32 columns")
        unknown = [name for name in selected if name not in available]
        if unknown:
            raise SourceError(f"unknown columns: {', '.join(unknown)}")

        profiles = []
        for name in selected:
            identifier = quote_identifier(name)
            numeric = is_numeric_type(str(available[name]))
            mean_expr = f", AVG({identifier})" if numeric else ""
            sql = (
                f"SELECT COUNT(*), COUNT({identifier}), COUNT(DISTINCT {identifier}), "
                f"MIN({identifier}), MAX({identifier}){mean_expr} "
                f"FROM {quote_identifier(relation)}"
            )
            row = self._internal_query(source, sql)[0]
            total, non_null, distinct, minimum, maximum, *mean = row
            profiles.append(
                {
                    "name": name,
                    "type": available[name],
                    "null_count": int(total) - int(non_null),
                    "distinct_count": int(distinct),
                    "min": to_json_value(minimum),
                    "max": to_json_value(maximum),
                    "mean": to_json_value(mean[0]) if mean else None,
                }
            )
        return {"source_id": source_id, "relation": relation, "columns": profiles}

    def sample_rows(self, source_id: str, relation: str, limit: int = 5) -> dict[str, Any]:
        source = self._source(source_id)
        self._require_relation(source, relation)
        safe_limit = min(max(int(limit), 1), 50)
        columns, rows = self._internal_query_with_columns(
            source,
            f"SELECT * FROM {quote_identifier(relation)} LIMIT {safe_limit}",
        )
        return {
            "source_id": source_id,
            "relation": relation,
            "columns": columns,
            "rows": [to_json_value(row) for row in rows],
            "row_count": len(rows),
        }

    def execute_sql(
        self,
        source_id: str,
        sql: str,
        max_rows: int | None = None,
    ) -> dict[str, Any]:
        source = self._source(source_id)
        limit = self.default_max_rows if max_rows is None else self._bounded_rows(max_rows)
        return self._execute_query(source, sql, sql, [], limit)

    def execute_analysis(self, source_id: str, spec: AnalysisSpec) -> dict[str, Any]:
        source = self._source(source_id)
        sql, parameters = self._build_analysis_sql(source, spec)
        display_sql = bind_sql_parameters(sql, parameters)
        result = self._execute_query(source, sql, display_sql, parameters, spec.limit)
        return {**result, "analysis_spec": spec.model_dump(mode="json")}

    def diagnose_table(
        self, source_id: str, relation: str,
        key_columns: list[str] | None = None, metric_columns: list[str] | None = None,
    ) -> dict[str, Any]:
        from .diagnostics import diagnose_table

        return diagnose_table(self, source_id, relation, key_columns, metric_columns)

    def compare_tables(self, spec: ComparisonSpec) -> dict[str, Any]:
        from .diagnostics import compare_tables

        return compare_tables(self, spec)

    def get_query_result(self, query_id: str) -> dict[str, Any]:
        with self._lock:
            record = self._queries.get(query_id)
            verified = query_id in self._verified_queries
        if record is None:
            raise SourceError(f"unknown query_id: {query_id}")
        source_current = self._result_source_is_current(record.result)
        if not source_current:
            with self._lock:
                self._verified_queries.discard(query_id)
        return {
            **record.result.model_dump(mode="json"),
            "verified": verified and source_current,
            "source_current": source_current,
        }

    def _execute_query(
        self,
        source: SourceInfo,
        executable_sql: str,
        display_sql: str,
        parameters: Sequence[Any],
        limit: int,
    ) -> dict[str, Any]:
        dialect = "sqlite" if source.kind == "sqlite" else "duckdb"
        validated = validate_read_only_sql(display_sql, dialect=dialect)
        limit = self._bounded_rows(limit)
        started = time.perf_counter()
        columns, fetched = self._query(source, executable_sql, limit + 1, parameters)
        elapsed_ms = max(0, round((time.perf_counter() - started) * 1_000))
        truncated = len(fetched) > limit
        rows = [to_json_value(row) for row in fetched[:limit]]
        result_digest = "sha256:" + self._digest({"columns": columns, "rows": rows})
        query_id = "qry_" + self._digest(
            {
                "source_id": source.source_id,
                "source_fingerprint": source.fingerprint,
                "sql": validated.normalized,
                "result_digest": result_digest,
            }
        )[:20]
        result = QueryResult(
            query_id=query_id,
            source_id=source.source_id,
            source_fingerprint=source.fingerprint,
            sql=validated.normalized,
            columns=columns,
            rows=rows,
            row_count=len(rows),
            truncated=truncated,
            elapsed_ms=elapsed_ms,
            result_digest=result_digest,
        )
        with self._lock:
            self._queries[query_id] = QueryRecord(result=result)
            self._verified_queries.discard(query_id)
        return result.model_dump(mode="json")

    def verify_query(self, query_id: str) -> dict[str, Any]:
        with self._lock:
            record = self._queries.get(query_id)
        if record is None:
            raise SourceError(f"unknown query_id: {query_id}")
        result = record.result
        warnings = []
        if result.row_count == 0:
            warnings.append("query returned no rows")
        if result.truncated:
            warnings.append("result was truncated; aggregate before drawing complete-set conclusions")
        source_current = self._result_source_is_current(result)
        if not source_current:
            warnings.append(
                "source file changed after this query was executed; run the analysis again"
            )
        checks = {
            "read_only": record.read_only,
            "single_statement": record.statement_count == 1,
            "stable_result_digest": result.result_digest.startswith("sha256:"),
            "column_row_shape": all(len(row) == len(result.columns) for row in result.rows),
            "source_unchanged": source_current,
        }
        valid = all(checks.values())
        if valid:
            with self._lock:
                self._verified_queries.add(query_id)
        return {
            "query_id": query_id,
            "valid": valid,
            "checks": checks,
            "warnings": warnings,
            "result_summary": {
                "source_id": result.source_id,
                "source_fingerprint": result.source_fingerprint,
                "sources": [source.model_dump(mode="json") for source in result.sources],
                "columns": result.columns,
                "row_count": result.row_count,
                "truncated": result.truncated,
                "result_digest": result.result_digest,
            },
        }

    def _result_source_is_current(self, result: QueryResult) -> bool:
        if result.sources:
            for item in result.sources:
                with self._lock:
                    source = self._sources.get(item.source_id)
                if source != item:
                    return False
                try:
                    self._source_path(item)
                except (OSError, ValueError):
                    return False
            return True
        with self._lock:
            source = self._sources.get(result.source_id)
        if source is None or source.fingerprint != result.source_fingerprint:
            return False
        try:
            self._source_path(source)
        except (OSError, ValueError):
            return False
        return True

    def _build_analysis_sql(
        self, source: SourceInfo, spec: AnalysisSpec
    ) -> tuple[str, list[Any]]:
        relation_names = [spec.relation]
        if spec.join is not None:
            relation_names.append(spec.join.relation)
        schemas: dict[str, dict[str, str]] = {}
        for relation in relation_names:
            described = self.describe_relation(source.source_id, relation)
            schemas[relation] = {
                str(column["name"]): str(column["type"])
                for column in described["columns"]
            }

        aliases = {spec.relation: "t0"}
        if spec.join is not None:
            if spec.join.relation == spec.relation:
                raise SourceError("join relation must differ from the primary relation")
            aliases[spec.join.relation] = "t1"
            self._require_field(spec.join.left, schemas)
            self._require_field(spec.join.right, schemas)
            if spec.join.left.relation != spec.relation:
                raise SourceError("join left field must belong to the primary relation")
            if spec.join.right.relation != spec.join.relation:
                raise SourceError("join right field must belong to the joined relation")
            right = self._field_sql(spec.join.right, aliases)
            counts = self._internal_query(
                source,
                f"SELECT COUNT({right}), COUNT(DISTINCT {right}) "
                f"FROM {quote_identifier(spec.join.relation)} AS t1",
            )[0]
            if int(counts[0]) != int(counts[1]):
                raise SourceError("joined relation key must be unique for analysis")

        selections: list[str] = []
        groups: list[str] = []
        output_names: set[str] = set()
        for dimension in spec.dimensions:
            self._require_field(dimension.field, schemas)
            expression = self._field_sql(dimension.field, aliases)
            if dimension.date_grain is not None:
                expression = date_bucket_sql(expression, dimension.date_grain, source.kind)
            alias = dimension.alias or dimension.field.column
            self._add_output_name(alias, output_names)
            selections.append(f"{expression} AS {quote_identifier(alias)}")
            groups.append(expression)

        for metric in spec.metrics:
            if metric.field is None:
                if metric.aggregation != "count":
                    raise SourceError(f"{metric.aggregation} metric requires a field")
                expression = "COUNT(*)"
            else:
                self._require_field(metric.field, schemas)
                field = self._field_sql(metric.field, aliases)
                functions = {
                    "sum": "SUM",
                    "avg": "AVG",
                    "min": "MIN",
                    "max": "MAX",
                    "count": "COUNT",
                    "count_distinct": "COUNT(DISTINCT",
                }
                if metric.aggregation == "count_distinct":
                    expression = f"COUNT(DISTINCT {field})"
                else:
                    expression = f"{functions[metric.aggregation]}({field})"
            self._add_output_name(metric.alias, output_names)
            selections.append(f"{expression} AS {quote_identifier(metric.alias)}")

        sql = (
            f"SELECT {', '.join(selections)} "
            f"FROM {quote_identifier(spec.relation)} AS t0"
        )
        if spec.join is not None:
            join_keyword = "INNER JOIN" if spec.join.kind == "inner" else "LEFT JOIN"
            sql += (
                f" {join_keyword} {quote_identifier(spec.join.relation)} AS t1 ON "
                f"{self._field_sql(spec.join.left, aliases)} = "
                f"{self._field_sql(spec.join.right, aliases)}"
            )

        parameters: list[Any] = []
        predicates: list[str] = []
        operator_sql = {
            "eq": "=",
            "ne": "<>",
            "gt": ">",
            "gte": ">=",
            "lt": "<",
            "lte": "<=",
        }
        for filter_spec in spec.filters:
            self._require_field(filter_spec.field, schemas)
            field = self._field_sql(filter_spec.field, aliases)
            if filter_spec.operator in operator_sql:
                predicates.append(f"{field} {operator_sql[filter_spec.operator]} ?")
                parameters.append(filter_spec.value)
            elif filter_spec.operator == "contains":
                if not isinstance(filter_spec.value, str):
                    raise SourceError("contains filter requires a string value")
                escaped = (
                    filter_spec.value.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
                )
                predicates.append(f"CAST({field} AS VARCHAR) LIKE ? ESCAPE '\'")
                parameters.append(f"%{escaped}%")
            elif filter_spec.operator == "in":
                if not isinstance(filter_spec.value, list) or len(filter_spec.value) == 0:
                    raise SourceError("in filter requires a non-empty list")
                placeholders = ", ".join("?" for _ in filter_spec.value)
                predicates.append(f"{field} IN ({placeholders})")
                parameters.extend(filter_spec.value)
            elif filter_spec.operator == "is_null":
                predicates.append(f"{field} IS NULL")
            elif filter_spec.operator == "not_null":
                predicates.append(f"{field} IS NOT NULL")
            else:
                raise SourceError(f"unsupported filter operator: {filter_spec.operator}")
        if predicates:
            sql += " WHERE " + " AND ".join(predicates)
        if groups:
            sql += " GROUP BY " + ", ".join(groups)
        if spec.sort:
            clauses = []
            for sort in spec.sort:
                if sort.column not in output_names:
                    raise SourceError(f"sort column is not an output column: {sort.column}")
                clauses.append(f"{quote_identifier(sort.column)} {sort.direction.upper()}")
            sql += " ORDER BY " + ", ".join(clauses)
        sql += f" LIMIT {spec.limit}"
        return sql, parameters

    @staticmethod
    def _add_output_name(name: str, output_names: set[str]) -> None:
        if not name.strip():
            raise SourceError("output column name must not be blank")
        if name in output_names:
            raise SourceError(f"duplicate output column name: {name}")
        output_names.add(name)

    @staticmethod
    def _require_field(field: FieldRef, schemas: dict[str, dict[str, str]]) -> None:
        relation = schemas.get(field.relation)
        if relation is None:
            raise SourceError(f"field relation is not part of the analysis: {field.relation}")
        if field.column not in relation:
            raise SourceError(f"unknown field: {field.relation}.{field.column}")

    @staticmethod
    def _field_sql(field: FieldRef, aliases: dict[str, str]) -> str:
        alias = aliases.get(field.relation)
        if alias is None:
            raise SourceError(f"relation is not part of the analysis: {field.relation}")
        return f"{alias}.{quote_identifier(field.column)}"

    def _source(self, source_id: str) -> SourceInfo:
        with self._lock:
            source = self._sources.get(source_id)
        if source is None:
            raise SourceError(f"unknown source_id: {source_id}")
        return source

    def _require_relation(self, source: SourceInfo, relation: str) -> None:
        relations = self.list_relations(source.source_id)["relations"]
        if relation not in relations:
            raise SourceError(f"unknown relation for {source.source_id}: {relation}")

    def _query(
        self,
        source: SourceInfo,
        sql: str,
        fetch_limit: int,
        parameters: Sequence[Any] = (),
    ) -> tuple[list[str], list[Any]]:
        if source.kind == "sqlite":
            with closing(self._connect_sqlite(source)) as connection:
                deadline = time.monotonic() + self.timeout_seconds
                connection.set_progress_handler(lambda: int(time.monotonic() >= deadline), 1_000)
                try:
                    cursor = connection.execute(sql, tuple(parameters))
                    columns = [item[0] for item in cursor.description or []]
                    return columns, cursor.fetchmany(fetch_limit)
                except sqlite3.OperationalError as error:
                    if "interrupt" in str(error).lower():
                        raise QueryTimeoutError(
                            f"query exceeded {self.timeout_seconds:g} seconds"
                        ) from error
                    raise

        with closing(self._connect_duckdb(source)) as connection:
            timer = threading.Timer(self.timeout_seconds, connection.interrupt)
            timer.daemon = True
            timer.start()
            try:
                cursor = connection.execute(sql, list(parameters))
                columns = [item[0] for item in cursor.description or []]
                return columns, cursor.fetchmany(fetch_limit)
            except duckdb.Error as error:
                if "interrupt" in str(error).lower():
                    raise QueryTimeoutError(
                        f"query exceeded {self.timeout_seconds:g} seconds"
                    ) from error
                raise
            finally:
                timer.cancel()

    def _internal_query(self, source: SourceInfo, sql: str) -> list[Any]:
        return self._query(source, sql, self.HARD_MAX_ROWS)[1]

    def _internal_query_with_columns(
        self, source: SourceInfo, sql: str
    ) -> tuple[list[str], list[Any]]:
        return self._query(source, sql, self.HARD_MAX_ROWS)

    def _connect_sqlite(self, source: SourceInfo) -> sqlite3.Connection:
        path = self._source_path(source)
        uri = f"file:{quote(path.as_posix(), safe='/:')}?mode=ro"
        connection = sqlite3.connect(uri, uri=True)
        connection.execute("PRAGMA query_only = ON")
        return connection

    def _connect_duckdb(self, source: SourceInfo) -> duckdb.DuckDBPyConnection:
        path = self._source_path(source)
        if source.kind == "duckdb":
            connection = duckdb.connect(str(path), read_only=True)
            connection.execute("SET enable_external_access = false")
            return connection
        connection = duckdb.connect(":memory:")
        if source.kind == "csv":
            if source.selection is None:
                literal = quote_literal(str(path))
                connection.execute(
                    f"CREATE TABLE {quote_identifier(self.CSV_RELATION)} AS "
                    f"SELECT * FROM read_csv_auto({literal}, sample_size = -1)"
                )
            else:
                self._create_sheet_table(connection, self.CSV_RELATION, self._read_csv(path, source.selection))
        elif source.kind == "xlsx":
            sheets, warnings = self._read_xlsx(path, source.selection)
            with self._lock:
                self._source_warnings[source.source_id] = warnings
            for name, sheet in sheets.items():
                self._create_sheet_table(connection, name, sheet)
        else:
            connection.close()
            raise SourceError(f"unsupported DuckDB-backed source kind: {source.kind}")
        connection.execute("SET enable_external_access = false")
        return connection

    def _source_path(self, source: SourceInfo) -> Path:
        """Re-check the path on every open to prevent symlink-swap escapes."""
        resolved = resolve_workspace_path(self.workspace_root, source.path)
        if workspace_relative_path(self.workspace_root, resolved) != source.path:
            raise SourceError("registered source path changed after registration")
        if self._file_fingerprint(resolved) != source.fingerprint:
            raise SourceError("registered source file changed after registration")
        return resolved

    @staticmethod
    def _create_sheet_table(connection: duckdb.DuckDBPyConnection, name: str, sheet: dict[str, Any]) -> None:
        definitions = ", ".join(
            f"{quote_identifier(column)} {type_name}"
            for column, type_name in zip(sheet["columns"], sheet["types"], strict=True)
        )
        connection.execute(f"CREATE TABLE {quote_identifier(name)} ({definitions})")
        if sheet["rows"]:
            placeholders = ", ".join("?" for _ in sheet["columns"])
            connection.executemany(
                f"INSERT INTO {quote_identifier(name)} VALUES ({placeholders})", sheet["rows"]
            )

    def _read_csv(self, path: Path, selection: TableSelection) -> dict[str, Any]:
        with path.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.reader(stream)
            header = next(islice(reader, selection.header_row - 1, None), None)
            if header is None:
                raise SourceError("CSV header row is absent")
            columns = normalize_headers(path.name, header)
            rows: list[list[Any]] = []
            for line, values in enumerate(reader, selection.header_row + 1):
                if line < (selection.data_start_row or selection.header_row + 1):
                    continue
                if selection.data_end_row is not None and line > selection.data_end_row:
                    break
                if len(values) > len(columns) and any(value.strip() for value in values[len(columns):]):
                    raise SourceError(f"CSV row {line} has values beyond the header")
                row = [value if value.strip() else None for value in values[:len(columns)]]
                row.extend([None] * (len(columns) - len(row)))
                if any(value is not None for value in row):
                    rows.append(row)
                if len(rows) * len(columns) > self.max_xlsx_cells:
                    raise SourceError("selected CSV data exceeds the configured cell limit")
        types = []
        for index in range(len(columns)):
            values = [row[index] for row in rows if row[index] is not None]
            if values and all(re.fullmatch(r"[+-]?(?:0|[1-9][0-9]*)(?:\.[0-9]+)?", value) for value in values):
                numeric_type = "DOUBLE" if any("." in value for value in values) else "BIGINT"
                types.append(numeric_type)
                for row in rows:
                    if row[index] is not None:
                        row[index] = float(row[index]) if numeric_type == "DOUBLE" else int(row[index])
            else:
                types.append("VARCHAR")
        return {"columns": columns, "types": types, "rows": rows}

    def _read_xlsx(
        self, path: Path, selection: TableSelection | None = None
    ) -> tuple[dict[str, dict[str, Any]], list[str]]:
        try:
            values_book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            formulas_book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        except Exception as error:
            raise SourceError(f"cannot read Excel workbook {path.name}: {error}") from error
        sheets: dict[str, dict[str, Any]] = {}
        warnings: list[str] = []
        consumed_cells = 0
        try:
            selected_names = [selection.sheet] if selection is not None and selection.sheet is not None else values_book.sheetnames
            for sheet_name in selected_names:
                if sheet_name not in values_book.sheetnames:
                    raise SourceError(f"unknown sheet: {sheet_name}")
                values_sheet = values_book[sheet_name]
                formulas_sheet = formulas_book[sheet_name]
                header_row = selection.header_row if selection is not None else 1
                data_start = (selection.data_start_row or header_row + 1) if selection is not None else 2
                data_end = selection.data_end_row if selection is not None else None
                value_rows = values_sheet.iter_rows(min_row=header_row, max_row=data_end, values_only=True)
                formula_rows = formulas_sheet.iter_rows(min_row=header_row, max_row=data_end, values_only=True)
                header = next(value_rows, None)
                formula_header = next(formula_rows, None)
                if header is None or formula_header is None:
                    raise SourceError(f"Excel sheet is empty: {sheet_name}")
                columns = normalize_headers(sheet_name, header)
                consumed_cells += len(columns)
                rows: list[list[Any]] = []
                missing_formula_cache = 0
                for line, (values, formulas) in enumerate(zip_longest(value_rows, formula_rows, fillvalue=()), header_row + 1):
                    if line < data_start:
                        continue
                    consumed_cells += len(columns)
                    if consumed_cells > self.max_xlsx_cells:
                        raise SourceError(
                            f"selected Excel sheets exceed {self.max_xlsx_cells} cell limit"
                        )
                    value_row = list(values[: len(columns)]) + [None] * max(0, len(columns) - len(values))
                    formula_row = list(formulas[: len(columns)]) + [None] * max(0, len(columns) - len(formulas))
                    if all(value is None for value in value_row) and all(
                        value is None for value in formula_row
                    ):
                        continue
                    for value, formula in zip(value_row, formula_row, strict=True):
                        if isinstance(formula, str) and formula.startswith("=") and value is None:
                            missing_formula_cache += 1
                    rows.append(value_row)
                types, mixed_columns = infer_column_types(columns, rows)
                coerced = [
                    [coerce_excel_value(value, type_name) for value, type_name in zip(row, types, strict=True)]
                    for row in rows
                ]
                if missing_formula_cache:
                    warnings.append(
                        f"{sheet_name}: {missing_formula_cache} formula cells have no cached value"
                    )
                if mixed_columns:
                    warnings.append(
                        f"{sheet_name}: mixed-type columns were converted to text: {', '.join(mixed_columns)}"
                    )
                sheets[sheet_name] = {"columns": columns, "types": types, "rows": coerced}
        finally:
            values_book.close()
            formulas_book.close()
        return sheets, warnings

    @staticmethod
    def _validate_kind(path: Path, kind: SourceKind) -> SourceKind:
        extensions = {
            "csv": {".csv"},
            "xlsx": {".xlsx"},
            "sqlite": {".sqlite", ".sqlite3", ".db"},
            "duckdb": {".duckdb"},
        }
        if kind not in extensions:
            raise SourceError(f"unsupported source kind: {kind}")
        if path.suffix.lower() not in extensions[kind]:
            raise SourceError(f"file extension does not match kind {kind}: {path.name}")
        return kind

    @classmethod
    def _bounded_rows(cls, value: int) -> int:
        value = int(value)
        if value <= 0:
            raise ValueError("max_rows must be positive")
        return min(value, cls.HARD_MAX_ROWS)

    @staticmethod
    def _digest(value: Any) -> str:
        return hashlib.sha256(canonical_json(value).encode("utf-8")).hexdigest()

    @staticmethod
    def _file_fingerprint(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
        return "sha256:" + digest.hexdigest()


def quote_identifier(identifier: str) -> str:
    if not identifier or "\x00" in identifier:
        raise SourceError("invalid SQL identifier")
    return '"' + identifier.replace('"', '""') + '"'


def quote_literal(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def normalize_headers(sheet_name: str, header: Sequence[Any]) -> list[str]:
    columns = ["" if value is None else str(value).strip() for value in header]
    if not columns or any(not column for column in columns):
        raise SourceError(f"Excel sheet has blank header cells: {sheet_name}")
    folded = [column.casefold() for column in columns]
    if len(set(folded)) != len(folded):
        raise SourceError(f"Excel sheet has duplicate header cells: {sheet_name}")
    return columns


def infer_column_types(columns: Sequence[str], rows: Sequence[Sequence[Any]]) -> tuple[list[str], list[str]]:
    types: list[str] = []
    mixed: list[str] = []
    for index, column in enumerate(columns):
        values = [row[index] for row in rows if row[index] is not None]
        kinds = {excel_value_kind(value) for value in values}
        if not kinds:
            type_name = "VARCHAR"
        elif kinds <= {"integer"}:
            type_name = "BIGINT"
        elif kinds <= {"integer", "number"}:
            type_name = "DOUBLE"
        elif kinds <= {"boolean"}:
            type_name = "BOOLEAN"
        elif kinds <= {"date"}:
            type_name = "DATE"
        elif kinds <= {"date", "datetime"}:
            type_name = "TIMESTAMP"
        elif kinds <= {"text"}:
            type_name = "VARCHAR"
        else:
            type_name = "VARCHAR"
            mixed.append(column)
        types.append(type_name)
    return types, mixed


def excel_value_kind(value: Any) -> str:
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, (float, Decimal)):
        return "number"
    if isinstance(value, datetime):
        return "datetime"
    if isinstance(value, date):
        return "date"
    return "text"


def coerce_excel_value(value: Any, type_name: str) -> Any:
    if value is None:
        return None
    if type_name == "VARCHAR":
        return value if isinstance(value, str) else str(value)
    if type_name == "DOUBLE":
        return float(value)
    if type_name == "TIMESTAMP" and isinstance(value, date) and not isinstance(value, datetime):
        return datetime.combine(value, datetime.min.time())
    return value


def date_bucket_sql(expression: str, grain: str, kind: SourceKind) -> str:
    if kind == "sqlite":
        formats = {"day": "%Y-%m-%d", "month": "%Y-%m-01", "year": "%Y-01-01"}
        return f"strftime('{formats[grain]}', {expression})"
    return f"CAST(date_trunc('{grain}', {expression}) AS DATE)"


def bind_sql_parameters(sql: str, parameters: Sequence[Any]) -> str:
    parts = sql.split("?")
    if len(parts) - 1 != len(parameters):
        raise SourceError("analysis parameter count does not match generated SQL")
    rendered = [parts[0]]
    for value, suffix in zip(parameters, parts[1:], strict=True):
        rendered.append(sql_literal(value))
        rendered.append(suffix)
    return "".join(rendered)


def sql_literal(value: Any) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return str(value)
    if isinstance(value, (date, datetime)):
        return quote_literal(value.isoformat())
    return quote_literal(str(value))


def is_numeric_type(type_name: str) -> bool:
    return bool(
        re.search(
            r"\b(TINYINT|SMALLINT|INTEGER|BIGINT|HUGEINT|UTINYINT|USMALLINT|UINTEGER|UBIGINT|REAL|FLOAT|DOUBLE|DECIMAL|NUMERIC)\b",
            type_name.upper(),
        )
    )

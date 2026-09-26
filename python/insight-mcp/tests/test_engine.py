from __future__ import annotations

from pathlib import Path

import duckdb
import pytest
from openpyxl import Workbook, load_workbook

from insight_mcp.engine import DataEngine
from insight_mcp.models import AnalysisSpec
from insight_mcp.security import PathSecurityError
from insight_mcp.sql_policy import SqlPolicyError


def test_csv_discovery_profile_query_and_verify(workspace: Path) -> None:
    engine = DataEngine(workspace)
    source = engine.register_source("sales.csv", "csv")
    source_id = source["source_id"]

    assert engine.list_relations(source_id)["relations"] == ["data"]
    described = engine.describe_relation(source_id, "data")
    assert [column["name"] for column in described["columns"]] == [
        "region",
        "amount",
        "sold_at",
        "note",
    ]
    profile = engine.profile_relation(source_id, "data", ["region", "amount"])
    assert profile["columns"][0]["distinct_count"] == 2

    result = engine.execute_sql(
        source_id,
        "SELECT region, SUM(amount) AS total FROM data GROUP BY region ORDER BY region",
    )
    assert result["rows"] == [["East", 22.0], ["West", 20.0]]
    assert engine.verify_query(result["query_id"])["valid"] is True


def test_sqlite_join_null_and_read_only(workspace: Path) -> None:
    engine = DataEngine(workspace)
    source_id = engine.register_source("shop.sqlite", "sqlite")["source_id"]
    assert engine.list_relations(source_id)["relations"] == ["customers", "orders"]

    result = engine.execute_sql(
        source_id,
        """
        SELECT c.name, COUNT(o.id) AS orders, SUM(o.amount) AS total
        FROM customers c LEFT JOIN orders o ON o.customer_id = c.id
        GROUP BY c.name ORDER BY c.name
        """,
    )
    assert result["rows"] == [["Ada", 2, 12.5], ["Lin", 1, 20.0]]

    with pytest.raises(SqlPolicyError):
        engine.execute_sql(source_id, "DELETE FROM orders")


def test_duckdb_discovery_sample_and_aggregation(workspace: Path) -> None:
    database = workspace / "metrics.duckdb"
    connection = duckdb.connect(str(database))
    connection.execute("CREATE TABLE metrics(category VARCHAR, value INTEGER)")
    connection.execute("INSERT INTO metrics VALUES ('a', 2), ('a', 3), ('b', NULL)")
    connection.close()

    engine = DataEngine(workspace)
    source_id = engine.register_source("metrics.duckdb", "duckdb")["source_id"]
    assert engine.list_relations(source_id)["relations"] == ["metrics"]
    assert engine.describe_relation(source_id, "metrics")["columns"][0]["name"] == "category"
    assert engine.sample_rows(source_id, "metrics", 2)["row_count"] == 2
    result = engine.execute_sql(
        source_id,
        "SELECT category, SUM(value) AS total FROM metrics GROUP BY category ORDER BY category",
    )
    assert result["rows"] == [["a", 5], ["b", None]]


def test_xlsx_structured_analysis_join_and_verification(workspace: Path) -> None:
    engine = DataEngine(workspace)
    source = engine.register_source("sales.xlsx", "xlsx")
    source_id = source["source_id"]
    assert engine.list_relations(source_id)["relations"] == ["客户", "订单"]
    assert source["warnings"] == ["订单: 3 formula cells have no cached value"]

    result = engine.execute_analysis(
        source_id,
        AnalysisSpec.model_validate(
            {
                "relation": "订单",
                "join": {
                    "relation": "客户",
                    "kind": "left",
                    "left": {"relation": "订单", "column": "customer_id"},
                    "right": {"relation": "客户", "column": "customer_id"},
                },
                "dimensions": [
                    {"field": {"relation": "客户", "column": "region"}}
                ],
                "metrics": [
                    {
                        "aggregation": "sum",
                        "field": {"relation": "订单", "column": "amount"},
                        "alias": "revenue",
                    }
                ],
                "sort": [{"column": "revenue", "direction": "desc"}],
            }
        ),
    )
    assert result["rows"] == [["East", 20.0], ["West", 20.0]]
    assert result["source_fingerprint"].startswith("sha256:")
    assert engine.get_query_result(result["query_id"])["verified"] is False
    assert engine.verify_query(result["query_id"])["valid"] is True
    assert engine.get_query_result(result["query_id"])["verified"] is True


def test_structured_filter_is_bound_and_source_changes_are_rejected(workspace: Path) -> None:
    engine = DataEngine(workspace)
    source_id = engine.register_source("sales.csv", "csv")["source_id"]
    result = engine.execute_analysis(
        source_id,
        AnalysisSpec.model_validate(
            {
                "relation": "data",
                "metrics": [{"aggregation": "count", "alias": "rows"}],
                "filters": [
                    {
                        "field": {"relation": "data", "column": "region"},
                        "operator": "eq",
                        "value": "East' OR 1=1 --",
                    }
                ],
            }
        ),
    )
    assert result["rows"] == [[0]]

    path = workspace / "sales.csv"
    path.write_text(path.read_text(encoding="utf-8") + "North,5,2025-01-04,ok\n", encoding="utf-8")
    with pytest.raises(ValueError, match="changed after registration"):
        engine.list_relations(source_id)
    verification = engine.verify_query(result["query_id"])
    assert verification["valid"] is False
    assert verification["checks"]["source_unchanged"] is False
    stale = engine.get_query_result(result["query_id"])
    assert stale["verified"] is False
    assert stale["source_current"] is False


def test_structured_date_grouping_and_contains_filter(workspace: Path) -> None:
    engine = DataEngine(workspace)
    source_id = engine.register_source("sales.csv", "csv")["source_id"]
    result = engine.execute_analysis(
        source_id,
        AnalysisSpec.model_validate(
            {
                "relation": "data",
                "dimensions": [
                    {
                        "field": {"relation": "data", "column": "sold_at"},
                        "date_grain": "month",
                        "alias": "month",
                    }
                ],
                "metrics": [
                    {
                        "aggregation": "sum",
                        "field": {"relation": "data", "column": "amount"},
                        "alias": "revenue",
                    }
                ],
                "filters": [
                    {
                        "field": {"relation": "data", "column": "region"},
                        "operator": "contains",
                        "value": "as",
                    }
                ],
            }
        ),
    )
    assert result["rows"] == [["2025-01-01", 22.0]]

    wildcard = engine.execute_analysis(
        source_id,
        AnalysisSpec.model_validate(
            {
                "relation": "data",
                "metrics": [{"aggregation": "count", "alias": "rows"}],
                "filters": [
                    {
                        "field": {"relation": "data", "column": "region"},
                        "operator": "contains",
                        "value": "%_",
                    }
                ],
            }
        ),
    )
    assert wildcard["rows"] == [[0]]


def test_xlsx_rejects_duplicate_join_key_and_import_limit(workspace: Path) -> None:
    duplicate_path = workspace / "duplicate.xlsx"
    workbook = load_workbook(workspace / "sales.xlsx")
    workbook["客户"].append([1, "Duplicate", "North"])
    workbook.save(duplicate_path)
    workbook.close()

    engine = DataEngine(workspace)
    source_id = engine.register_source("duplicate.xlsx", "xlsx")["source_id"]
    spec = AnalysisSpec.model_validate(
        {
            "relation": "订单",
            "join": {
                "relation": "客户",
                "left": {"relation": "订单", "column": "customer_id"},
                "right": {"relation": "客户", "column": "customer_id"},
            },
            "metrics": [{"aggregation": "count", "alias": "orders"}],
        }
    )
    with pytest.raises(ValueError, match="key must be unique"):
        engine.execute_analysis(source_id, spec)

    limited = DataEngine(workspace, max_xlsx_cells=5)
    with pytest.raises(ValueError, match="cell limit"):
        limited.register_source("sales.xlsx", "xlsx")


def test_xlsx_rejects_blank_and_duplicate_headers(workspace: Path) -> None:
    path = workspace / "bad-headers.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["region", "", "region"])
    sheet.append(["East", 1, "duplicate"])
    workbook.save(path)
    workbook.close()

    engine = DataEngine(workspace)
    with pytest.raises(ValueError, match="blank header"):
        engine.register_source("bad-headers.xlsx", "xlsx")


def test_query_result_is_runtime_scoped(workspace: Path) -> None:
    first = DataEngine(workspace)
    source_id = first.register_source("sales.csv", "csv")["source_id"]
    query_id = first.execute_sql(source_id, "SELECT COUNT(*) AS rows FROM data")["query_id"]
    first.verify_query(query_id)

    restarted = DataEngine(workspace)
    with pytest.raises(ValueError, match="unknown query_id"):
        restarted.get_query_result(query_id)


def test_result_truncation_and_stable_query_id(workspace: Path) -> None:
    engine = DataEngine(workspace, default_max_rows=1)
    source_id = engine.register_source("sales.csv", "csv")["source_id"]
    first = engine.execute_sql(source_id, "SELECT region FROM data ORDER BY region")
    second = engine.execute_sql(source_id, "SELECT region FROM data ORDER BY region")
    assert first["truncated"] is True
    assert first["row_count"] == 1
    assert first["query_id"] == second["query_id"]
    assert engine.verify_query(first["query_id"])["warnings"]


def test_rechecks_registered_path_before_each_open(
    workspace: Path, tmp_path_factory: pytest.TempPathFactory
) -> None:
    engine = DataEngine(workspace)
    source_id = engine.register_source("sales.csv", "csv")["source_id"]
    external = tmp_path_factory.mktemp("swap-external") / "sales.csv"
    external.write_text("secret\nvalue\n", encoding="utf-8")
    original = workspace / "sales.csv"
    original.unlink()
    try:
        original.symlink_to(external)
    except OSError:
        pytest.skip("symlink creation is unavailable on this platform")
    with pytest.raises(PathSecurityError, match="inside"):
        engine.list_relations(source_id)
